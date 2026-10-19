import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Kế hoạch Tháng 10-2026"

headers = ["Tên công việc", "Người phụ trách", "Trạng thái", "Mức độ ưu tiên", "Hạn chót", "Vấn đề gặp phải (Blocker)", "Ghi chú"]
ws.append(headers)

# Styling header
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

sample_data = [
    ["Tối ưu hóa Database Query và Indexing", "Dev An", "TODO", "HIGH", "2026-10-25", "", "Tối ưu hoá các truy vấn bảng tasks và daily notes"],
    ["Thiết kế luồng Export Báo cáo PDF/Excel", "Dev Bình", "IN_PROGRESS", "MEDIUM", "2026-10-28", "", "Hỗ trợ xuất báo cáo tiến độ sprint cho Tech Lead"],
    ["Tích hợp Payment Gateway VNPay Sandbox", "Dev Cường", "BLOCKED", "HIGH", "2026-10-24", "Đang chờ đối tác VNPay cấp Merchant Key Sandbox", "Cần PM liên hệ hỗ trợ thúc đẩy"],
    ["Kiểm thử bảo mật Pentest hệ thống API", "Dev QA", "TODO", "HIGH", "2026-10-30", "", "Kiểm tra lỗ hổng OWASP Top 10 và JWT token"],
    ["Lên lịch họp Sprint Review với Khách hàng", "PM Tech Lead", "TODO", "URGENT", "2026-10-27", "", "Gửi email báo cáo tiến độ và demo các tính năng mới"]
]

for row in sample_data:
    ws.append(row)

# Adjust column widths
for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

wb.save("data/monthly_report_sample.xlsx")
print("SUCCESS: Sample Excel created at data/monthly_report_sample.xlsx")
